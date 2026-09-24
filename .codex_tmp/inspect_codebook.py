from pathlib import Path
import json
from docx import Document
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
doc_path = root / "Codebook" / "Codebook.docx"
xlsx_doc_path = root / "Codebook" / "Output" / "Codebook_final.docx"
new_codes_path = root / "Codebook" / "Output" / "New Codes.docx"
xlsx_path = root / "Codebook" / "Output" / "Codebook_final.xlsx"

doc = Document(doc_path)
doc_data = {
    "paragraphs": [
        {"index": i, "style": p.style.name if p.style else None, "text": p.text}
        for i, p in enumerate(doc.paragraphs)
        if p.text.strip()
    ],
    "tables": [],
}
for ti, table in enumerate(doc.tables):
    rows = []
    for ri, row in enumerate(table.rows):
        rows.append([cell.text for cell in row.cells])
    doc_data["tables"].append({"index": ti, "rows": rows})

wb = load_workbook(xlsx_path, data_only=False, read_only=False)
xlsx_data = {"sheets": []}
for ws in wb.worksheets:
    values = []
    for row in ws.iter_rows():
        vals = [cell.value for cell in row]
        if any(v is not None for v in vals):
            values.append({"row": row[0].row, "values": vals})
    xlsx_data["sheets"].append({
        "title": ws.title,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "rows": values,
    })

out_dir = root / ".codex_tmp"
out_dir.mkdir(exist_ok=True)
(out_dir / "codebook_doc.json").write_text(json.dumps(doc_data, ensure_ascii=False, indent=2), encoding="utf-8")

if xlsx_doc_path.exists():
    xdoc = Document(xlsx_doc_path)
    xdoc_data = {
        "paragraphs": [
            {"index": i, "style": p.style.name if p.style else None, "text": p.text}
            for i, p in enumerate(xdoc.paragraphs)
            if p.text.strip()
        ],
        "tables": [
            {"index": ti, "rows": [[cell.text for cell in row.cells] for row in table.rows]}
            for ti, table in enumerate(xdoc.tables)
        ],
    }
    (out_dir / "codebook_final_doc.json").write_text(json.dumps(xdoc_data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"final doc paragraphs={len(xdoc.paragraphs)} tables={len(xdoc.tables)}")

if new_codes_path.exists():
    ndoc = Document(new_codes_path)
    ndoc_data = {
        "paragraphs": [
            {"index": i, "style": p.style.name if p.style else None, "text": p.text}
            for i, p in enumerate(ndoc.paragraphs)
            if p.text.strip()
        ],
        "tables": [
            {"index": ti, "rows": [[cell.text for cell in row.cells] for row in table.rows]}
            for ti, table in enumerate(ndoc.tables)
        ],
    }
    (out_dir / "new_codes_doc.json").write_text(json.dumps(ndoc_data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"new codes doc paragraphs={len(ndoc.paragraphs)} tables={len(ndoc.tables)}")
(out_dir / "codebook_xlsx.json").write_text(json.dumps(xlsx_data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
print(f"doc paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")
for table in doc_data["tables"]:
    print(f"doc table {table['index']}: {len(table['rows'])} rows x {max((len(r) for r in table['rows']), default=0)} cols")
for sheet in xlsx_data["sheets"]:
    print(f"xlsx sheet {sheet['title']}: {sheet['max_row']} rows x {sheet['max_column']} cols; nonempty rows={len(sheet['rows'])}")
