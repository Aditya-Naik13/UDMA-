from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "Codebook" / "Output" / "Codebook_final.xlsx"

rows_by_tag = {
    "COMMUNITY_BUILT": [45, 51, 49, 52, 53, 48, 50, 46, 47],
    "CONSULT_NO_POWER": list(range(2, 12)),
    "PERF_VS_AUTH": [5],
    "NEGATIVE_SELFEFF": [7],
    "DIGNITY_STRUCT": [54, 55, 56, 57],
    "COGOVERN": [45, 49, 19],
    "THROUGHPUT": [21, 20, 23],
    "SAFETY_ENG": [22, 43, 15],
    "TRUST_METRIC": [13, 12, 14, 15, 19],
    "EXPERTS_CENTERED": [30, 31, 32, 33, 34, 36, 38, 39],
    "SURVEILLANCE": [43],
    "FRICTION_MANAGED": [43, 40],
    "CHARITY_REFRAIM": [54, 55],
    "CONDITIONALITY": [55],
    "ACCESS_AS_EFF": [28, 29],
    "APARTHEID_GEO": [29, 39, 51],
    "EFFICIENCY_SURPLUS": [29],
}

wb = load_workbook(BOOK, data_only=True, read_only=True)
ws = wb["Codebook"]
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
records = {}
for excel_row, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    record = dict(zip(headers, values))
    if record.get("Tag") in rows_by_tag:
        record["_row"] = excel_row
        records[record["Tag"]] = record

lws = wb["LitEvidenceMap"]
for tag, lit_rows in rows_by_tag.items():
    r = records[tag]
    papers = []
    for row in lit_rows:
        papers.append(f"{lws.cell(row, 6).value} {lws.cell(row, 7).value}")
    example = str(r.get("Example quote") or "")
    has_verbatim = bool(r.get("Locator")) and not example.startswith("[")
    print(
        f"{tag}\tCodebook!{r['_row']}\tpapers={len(lit_rows)}\t"
        f"locator={'yes' if r.get('Locator') else 'no'}\tverbatim={'yes' if has_verbatim else 'no'}\t"
        f"{'; '.join(papers)}"
    )
