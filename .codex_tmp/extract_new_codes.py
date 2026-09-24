from pathlib import Path
import json
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "Codebook" / "Output" / "Codebook_final.xlsx"
TAGS = [
    "COMMUNITY_BUILT", "CONSULT_NO_POWER", "PERF_VS_AUTH", "NEGATIVE_SELFEFF",
    "DIGNITY_STRUCT", "COGOVERN", "THROUGHPUT", "SAFETY_ENG", "TRUST_METRIC",
    "EXPERTS_CENTERED", "SURVEILLANCE", "FRICTION_MANAGED", "CHARITY_REFRAIM",
    "CONDITIONALITY", "ACCESS_AS_EFF", "APARTHEID_GEO", "EFFICIENCY_SURPLUS",
]

wb = load_workbook(BOOK, data_only=True, read_only=True)
ws = wb["Codebook"]
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
codes = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    record = dict(zip(headers, row))
    if record.get("Tag") in TAGS:
        codes[record["Tag"]] = record

lws = wb["LitEvidenceMap"]
lheaders = [c.value for c in next(lws.iter_rows(min_row=1, max_row=1))]
papers = [dict(zip(lheaders, row)) for row in lws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in row)]

output = []
for tag in TAGS:
    rec = codes[tag]
    refs = [p for p in papers if tag in str(p.get("Feeds codes") or "")]
    output.append({"code": rec, "references": refs})

out = ROOT / ".codex_tmp" / "new_codes_source.json"
out.write_text(json.dumps(output, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
for item in output:
    print(f"{item['code']['Tag']}: {len(item['references'])} evidence-map rows")
