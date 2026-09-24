from pathlib import Path
from collections import OrderedDict
from copy import deepcopy

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Codebook" / "Output" / "Codebook_final.xlsx"
OUTPUT = ROOT / "Codebook" / "Output" / "New Codes Revised.docx"

TAGS = [
    "COMMUNITY_BUILT", "CONSULT_NO_POWER", "PERF_VS_AUTH", "NEGATIVE_SELFEFF",
    "DIGNITY_STRUCT", "COGOVERN", "THROUGHPUT", "SAFETY_ENG", "TRUST_METRIC",
    "EXPERTS_CENTERED", "SURVEILLANCE", "FRICTION_MANAGED", "CHARITY_REFRAIM",
    "CONDITIONALITY", "ACCESS_AS_EFF", "APARTHEID_GEO", "EFFICIENCY_SURPLUS",
]

# Exact LitEvidenceMap worksheet rows supporting the author-year list recorded in
# each Codebook row. Keeping these as workbook row references prevents adding
# papers that merely share a broader mechanism but are not named for the code.
CITATION_ROWS = {
    "COMMUNITY_BUILT": [45, 51, 49, 52, 53, 48, 50, 46, 47],
    "CONSULT_NO_POWER": list(range(2, 12)),
    "PERF_VS_AUTH": [5],
    "NEGATIVE_SELFEFF": [7],
    "DIGNITY_STRUCT": [54, 55, 56, 57],
    "COGOVERN": [45, 49, 19],
    "THROUGHPUT": [21, 20, 23],
    "SAFETY_ENG": [22, 43, 15],
    "TRUST_METRIC": [13, 12, 14, 15, 19],
    "EXPERTS_CENTERED": [30, 31, 32, 33, 34, 36, 38, 39],
    "SURVEILLANCE": [43],
    "FRICTION_MANAGED": [43, 40],
    "CHARITY_REFRAIM": [54, 55],
    "CONDITIONALITY": [55],
    "ACCESS_AS_EFF": [28, 29],
    "APARTHEID_GEO": [29, 39, 51],
    "EFFICIENCY_SURPLUS": [29],
}

SECTION_ORDER = [
    ("Codes Added to Existing Domains", ["C", "E", "G"]),
    ("Domain H The Discourse the AV Carries", ["H"]),
    ("Domain I The Charity Model and Digital Food System", ["I"]),
]

DOMAIN_HEADINGS = {
    "C": "Domain C Alignment and What Works",
    "E": "Domain E Participation Gap and Infrastructure Without Community",
    "G": "Domain G Mobility as Social Service Target State",
}

NAVY = "17365D"
PALE_BLUE = "EAF1F8"
PALE_GRAY = "F5F6F7"
BORDER = "D9D9D9"
TEXT = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="Arial", size=10.5, bold=False, italic=False, color=TEXT):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    r_pr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "19")
    r_pr.append(size)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])
    set_font(run, size=9, color=RGBColor(90, 90, 90))


def add_table_text(cell, text, bold=False, color=TEXT, size=10.2):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(str(text))
    set_font(run, size=size, bold=bold, color=color)


def value_or_na(value):
    if value is None or str(value).strip() == "":
        return "Not recorded"
    return str(value).strip()


def normalize_example(value):
    text = value_or_na(value)
    if text.startswith("[LIT]"):
        return "No transcript example is recorded. Use the literature evidence and references listed below."
    if text.startswith("[No verbatim quote"):
        return "No verbatim quote is recorded. The workbook points to design principle 11, “Build it with the community.”"
    return text


wb = load_workbook(SOURCE, data_only=True, read_only=False)
code_ws = wb["Codebook"]
code_headers = [c.value for c in code_ws[1]]
codes = OrderedDict()
for row in code_ws.iter_rows(min_row=2, values_only=True):
    rec = dict(zip(code_headers, row))
    if rec.get("Tag") in TAGS:
        codes[rec["Tag"]] = rec

if list(codes.keys()) != TAGS:
    missing = [tag for tag in TAGS if tag not in codes]
    raise RuntimeError(f"Missing code rows: {missing}; workbook order was {list(codes)}")

lit_ws = wb["LitEvidenceMap"]
lit_headers = [c.value for c in lit_ws[1]]
lit_by_row = {}
for row_num in range(2, lit_ws.max_row + 1):
    values = [lit_ws.cell(row=row_num, column=col).value for col in range(1, len(lit_headers) + 1)]
    if values[0] and values[0] not in {"Thread"} and not str(values[0]).startswith("Who is centered") and not str(values[0]).startswith("Thesis clause support"):
        lit_by_row[row_num] = dict(zip(lit_headers, values))

for tag, rows in CITATION_ROWS.items():
    for row_num in rows:
        if row_num not in lit_by_row:
            raise RuntimeError(f"{tag} points to missing LitEvidenceMap row {row_num}")

# Assign stable reference IDs in first-use order, deduplicated by DOI/URL or title.
reference_by_key = OrderedDict()
code_reference_ids = {}
for tag in TAGS:
    ids = []
    for row_num in CITATION_ROWS[tag]:
        ref = lit_by_row[row_num]
        key = value_or_na(ref.get("DOI or URL"))
        if key in {"Not found", "Not recorded"}:
            key = f"title::{ref.get('Paper title')}"
        if key not in reference_by_key:
            reference_by_key[key] = {"id": f"R{len(reference_by_key) + 1:02d}", "row": row_num, **ref}
        ids.append(reference_by_key[key]["id"])
    code_reference_ids[tag] = ids

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.5)
normal.font.color.rgb = TEXT
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.1

title_style = styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
title_style.font.size = Pt(24)
title_style.font.bold = True
title_style.font.color.rgb = TEXT
title_style.paragraph_format.space_after = Pt(9)
title_ppr = title_style._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for style_name, size, before, after in (
    ("Heading 1", 17, 15, 7),
    ("Heading 2", 13, 12, 5),
    ("Heading 3", 11.5, 9, 4),
):
    st = styles[style_name]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = TEXT
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run("New literature derived codes   ")
set_font(fp.runs[0], size=9, color=RGBColor(90, 90, 90))
add_page_field(fp)

title = doc.add_paragraph(style="Title")
title.add_run("New Literature Derived Codes")

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
r = subtitle.add_run("Mobility as Social Service research codebook")
set_font(r, size=11.5, italic=True, color=RGBColor(70, 70, 70))

p = doc.add_paragraph()
p.add_run(
    "The final codebook identifies 17 Tier L codes that were added from the research literature. "
    "Six extend existing domains C, E, and G; six form domain H; and five form domain I. "
    "This guide presents the coding boundaries and evidence recorded in the workbook so each code can be applied consistently."
)

p = doc.add_paragraph()
lead = p.add_run("Source of record. ")
set_font(lead, bold=True)
p.add_run(
    "Codebook_final.xlsx, especially the Codebook and LitEvidenceMap sheets. Definitions, inclusion and exclusion rules, examples, locators, crosswalks, and citations below come from those sheets."
)

h = doc.add_paragraph(style="Heading 1")
h.add_run("How to Use This Guide")

for text in (
    "Use Definition to identify the concept, Include when to decide what qualifies, and Exclude when to separate nearby codes.",
    "Tier L means the literature introduced the code. Some Tier L codes also have a transcript example or workshop cross-reference.",
    "Reference numbers connect each code to full paper records in the References section. A missing DOI or URL is stated explicitly rather than supplied from outside the workbook.",
):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)

h = doc.add_paragraph(style="Heading 1")
h.add_run("Section Map")

map_table = doc.add_table(rows=1, cols=3)
map_table.alignment = WD_TABLE_ALIGNMENT.CENTER
map_table.autofit = False
map_table.columns[0].width = Inches(2.1)
map_table.columns[1].width = Inches(0.8)
map_table.columns[2].width = Inches(4.0)
headers = ["Section", "Codes", "Domains"]
for i, text in enumerate(headers):
    cell = map_table.rows[0].cells[i]
    set_cell_shading(cell, NAVY)
    add_table_text(cell, text, bold=True, color=RGBColor(255, 255, 255))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
set_repeat_table_header(map_table.rows[0])
rows = [
    ("Existing domains", "6", "C Alignment and what works; E Participation gap; G Mobility as Social Service"),
    ("Domain H", "6", "The discourse the AV carries"),
    ("Domain I", "5", "The charity model and digital food system"),
]
for ri, row_data in enumerate(rows, start=1):
    cells = map_table.add_row().cells
    if ri % 2 == 0:
        for cell in cells:
            set_cell_shading(cell, PALE_GRAY)
    for i, text in enumerate(row_data):
        add_table_text(cells[i], text)
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cells[i])
    cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    prevent_row_split(map_table.rows[-1])
set_table_borders(map_table)


def add_code_record(rec, heading_style="Heading 2"):
    tag = rec["Tag"]
    heading = doc.add_paragraph(style=heading_style)
    heading.add_run(tag.replace("_", " "))

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.45)
    table.columns[1].width = Inches(5.55)
    set_table_borders(table)

    header_cells = table.rows[0].cells
    header_cells[0].merge(header_cells[1])
    header = header_cells[0]
    set_cell_shading(header, NAVY)
    meta = f"{tag}   Tier {value_or_na(rec['Tier'])}   Thesis clause {value_or_na(rec['Paper map (clause)'])}"
    add_table_text(header, meta, bold=True, color=RGBColor(255, 255, 255), size=10.5)
    set_cell_margins(header, top=105, bottom=105)
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])

    crosswalk_parts = [
        f"OG: {value_or_na(rec['OG source'])}",
        f"MaSS: {value_or_na(rec['MaSS source'])}",
        f"Capstone: {value_or_na(rec['Capstone source'])}",
    ]
    ref_labels = []
    for ref_id, row_num in zip(code_reference_ids[tag], CITATION_ROWS[tag]):
        ref = lit_by_row[row_num]
        ref_labels.append(f"[{ref_id}] {value_or_na(ref['Author(s)'])} {value_or_na(ref['Year'])}")
    if tag == "SURVEILLANCE":
        ref_labels.append("Capstone C04 face verification")

    detail_rows = [
        ("Domain", f"{rec['Domain']}  {value_or_na(rec['Prose title'])}"),
        ("Definition", value_or_na(rec["Definition"])),
        ("Include when", value_or_na(rec["Inclusion"])),
        ("Exclude when", value_or_na(rec["Exclusion"]).replace("Lite evidence map", "Lit evidence map")),
        ("Example", normalize_example(rec["Example quote"])),
        ("Locator", value_or_na(rec["Locator"])),
        ("Crosswalk", "\n".join(crosswalk_parts)),
        ("Literature basis", value_or_na(rec["Lit evidence"])),
        ("Cited sources", "; ".join(ref_labels)),
    ]

    for row_index, (label, value) in enumerate(detail_rows, start=1):
        cells = table.add_row().cells
        set_cell_shading(cells[0], PALE_BLUE)
        if row_index % 2 == 0:
            set_cell_shading(cells[1], PALE_GRAY)
        add_table_text(cells[0], label, bold=True)
        add_table_text(cells[1], value)
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
        prevent_row_split(table.rows[-1])

    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)


for section_title, domains in SECTION_ORDER:
    section_heading = doc.add_paragraph(style="Heading 1")
    section_heading.add_run(section_title)
    section_heading.paragraph_format.page_break_before = True
    for domain in domains:
        domain_records = [codes[tag] for tag in TAGS if codes[tag]["Domain"] == domain]
        if len(domains) > 1:
            d_heading = doc.add_paragraph(style="Heading 2")
            d_heading.add_run(DOMAIN_HEADINGS[domain])
            for rec in domain_records:
                add_code_record(rec, heading_style="Heading 3")
        else:
            for rec in domain_records:
                add_code_record(rec)

# Reference appendix
ref_heading = doc.add_paragraph(style="Heading 1")
ref_heading.add_run("References")
ref_heading.paragraph_format.page_break_before = True

p = doc.add_paragraph()
p.add_run(
    "Reference details below reproduce the LitEvidenceMap sheet. Journal and proceedings information is retained as recorded."
)

for ref in reference_by_key.values():
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.05
    lead = p.add_run(f"[{ref['id']}] ")
    set_font(lead, size=9.5, bold=True)
    authors = value_or_na(ref["Author(s)"]).rstrip().rstrip(".")
    paper_title = value_or_na(ref["Paper title"]).rstrip().rstrip(".")
    venue = value_or_na(ref["Journal or proceedings"]).rstrip().rstrip(".")
    citation = f"{authors}. ({value_or_na(ref['Year'])}). {paper_title}. {venue}. "
    body = p.add_run(citation)
    set_font(body, size=9.5)
    doi = value_or_na(ref["DOI or URL"])
    if doi.lower().startswith("http"):
        add_hyperlink(p, doi, doi)
    elif doi == "Not found":
        missing = p.add_run("DOI or URL not recorded in the evidence map.")
        set_font(missing, size=9.5, italic=True, color=RGBColor(90, 90, 90))
    else:
        # The workbook contains one bare DOI; convert only its display into a link
        # by adding the standard resolver prefix without changing the identifier.
        url = f"https://doi.org/{doi}"
        add_hyperlink(p, url, url)

# Improve pagination: keep headings and the next paragraph/table together.
for paragraph in doc.paragraphs:
    if paragraph.style and paragraph.style.name.startswith("Heading"):
        paragraph.paragraph_format.keep_with_next = True

doc.core_properties.title = "New Literature Derived Codes"
doc.core_properties.subject = "Clarified guide to the 17 literature-derived codes in the Mobility as Social Service codebook"
doc.core_properties.keywords = "codebook, mobility, social service, literature, qualitative coding"
doc.save(OUTPUT)
print(f"Saved {OUTPUT}")
print(f"Codes: {len(codes)}")
print(f"Unique references: {len(reference_by_key)}")
